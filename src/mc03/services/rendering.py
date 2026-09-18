"""Fixed-template workbook rendering for completed RCBC demo runs.

The renderer discovers approved ``Output_Mapping_Regions`` from workbook defined
names, tables, or explicit marker cells.  It writes only the discovered cells,
loads formulas as formulas, and commits both output workbooks only after both
renders have succeeded.
"""

from __future__ import annotations

import os
import tempfile
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from mc03.domain.models import ProcessedRow, RunResult

TARGET_SHEETS: tuple[str, ...] = ("FIELD RSULT", "DRR")
OUTPUT_MAPPING_MARKER = "output_mapping_regions"
# Deployments may supply approved coordinates explicitly.  Empty defaults force
# the renderer to discover regions from the fixed template metadata instead of
# silently writing to arbitrary cells.
OUTPUT_MAPPING_REGIONS: dict[str, tuple[str, ...]] = {}
ROW_VALUE_FIELDS: tuple[str, ...] = (
    "ch_code",
    "account_number",
    "relation",
    "csu_rank",
    "selected_rfd",
    "sanitized_remark",
    "remark_length",
    "disposition",
)


class RenderError(RuntimeError):
    """Raised when a fixed template cannot be safely rendered."""


@dataclass(frozen=True, slots=True)
class MappingRegion:
    """One rectangular writable region in a target worksheet."""

    sheet_name: str
    min_row: int
    min_col: int
    max_row: int
    max_col: int

    @property
    def width(self) -> int:
        """Return the number of writable columns."""
        return self.max_col - self.min_col + 1

    @property
    def height(self) -> int:
        """Return the number of writable rows."""
        return self.max_row - self.min_row + 1


def _safe_output_stem(run_id: str) -> str:
    """Return a filename-safe, non-secret output stem."""
    del run_id
    return f"rcbc_{uuid4().hex}"


def _sheet_name(value: object) -> str:
    """Normalize a worksheet name for target-sheet comparison."""
    return str(value).strip().casefold()


def _parse_destination(sheet_name: str, reference: str) -> MappingRegion | None:
    """Parse an openpyxl defined-name destination into a mapping region."""
    text = reference.strip()
    if "!" in text:
        destination_sheet, cell_range = text.rsplit("!", 1)
        destination_sheet = destination_sheet.strip("'\"")
    else:
        destination_sheet, cell_range = sheet_name, text
    cell_range = cell_range.replace("$", "")
    try:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    except ValueError:
        return None
    return MappingRegion(destination_sheet, min_row, min_col, max_row, max_col)


def _named_regions(workbook: Any) -> list[MappingRegion]:
    """Find regions declared through workbook-level defined names."""
    regions: list[MappingRegion] = []
    defined_names = getattr(workbook, "defined_names", {})
    values = defined_names.values() if hasattr(defined_names, "values") else ()
    for defined_name in values:
        name = str(getattr(defined_name, "name", ""))
        if OUTPUT_MAPPING_MARKER not in name.casefold().replace(" ", "_"):
            continue
        try:
            destinations = list(defined_name.destinations)
        except (AttributeError, KeyError, TypeError, ValueError):
            continue
        for sheet_name, reference in destinations:
            region = _parse_destination(str(sheet_name), str(reference))
            if region is not None:
                regions.append(region)
    return regions


def _table_regions(workbook: Any) -> list[MappingRegion]:
    """Find regions represented by tables named ``Output_Mapping_Regions``."""
    regions: list[MappingRegion] = []
    for worksheet in workbook.worksheets:
        tables = getattr(worksheet, "tables", {})
        values = tables.values() if hasattr(tables, "values") else ()
        for table in values:
            name = str(getattr(table, "displayName", getattr(table, "name", "")))
            if OUTPUT_MAPPING_MARKER not in name.casefold().replace(" ", "_"):
                continue
            parsed = _parse_destination(worksheet.title, str(table.ref))
            if parsed is not None:
                regions.append(parsed)
    return regions


def _marker_regions(workbook: Any) -> list[MappingRegion]:
    """Find marker-cell regions, supporting simple hand-built test templates."""
    regions: list[MappingRegion] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str):
                    continue
                normalized = value.strip().casefold().replace(" ", "_")
                if not normalized.startswith(OUTPUT_MAPPING_MARKER):
                    continue
                explicit = "=" in value
                if explicit:
                    _, reference = value.split("=", 1)
                    parsed = _parse_destination(worksheet.title, reference)
                    if parsed is not None:
                        regions.append(parsed)
                        continue
                min_row = cell.row + 1
                min_col = cell.column
                if min_row <= worksheet.max_row and min_col <= worksheet.max_column:
                    regions.append(
                        MappingRegion(
                            worksheet.title,
                            min_row,
                            min_col,
                            worksheet.max_row,
                            worksheet.max_column,
                        )
                    )
    return regions


def _configured_regions() -> list[MappingRegion]:
    """Return regions supplied by an approved application-level mapping."""
    regions: list[MappingRegion] = []
    for sheet_name, references in OUTPUT_MAPPING_REGIONS.items():
        values: Iterable[object] = (references,) if isinstance(references, str) else references
        for reference in values:
            parsed = _parse_destination(str(sheet_name), str(reference))
            if parsed is not None:
                regions.append(parsed)
    return regions


def _discover_regions(workbook: Any) -> dict[str, list[MappingRegion]]:
    """Discover and validate one or more mapped regions per required sheet."""
    candidates = (
        _configured_regions()
        + _named_regions(workbook)
        + _table_regions(workbook)
        + _marker_regions(workbook)
    )
    grouped: dict[str, list[MappingRegion]] = {name: [] for name in TARGET_SHEETS}
    for region in candidates:
        matching_sheet = next(
            (name for name in TARGET_SHEETS if _sheet_name(name) == _sheet_name(region.sheet_name)),
            None,
        )
        if matching_sheet is not None:
            normalized = MappingRegion(
                matching_sheet,
                region.min_row,
                region.min_col,
                region.max_row,
                region.max_col,
            )
            if normalized not in grouped[matching_sheet]:
                grouped[matching_sheet].append(normalized)
    missing = [sheet for sheet, regions in grouped.items() if not regions]
    if missing:
        raise RenderError(
            "Output_Mapping_Regions missing from sheet(s): " + ", ".join(missing)
        )
    return grouped


def _looks_like_header(worksheet: Any, region: MappingRegion) -> bool:
    """Detect a header row so it remains untouched inside a table-like region."""
    values = [
        worksheet.cell(region.min_row, column).value
        for column in range(region.min_col, region.max_col + 1)
    ]
    normalized = {_sheet_name(value).replace(" ", "_") for value in values if value is not None}
    return bool(
        normalized
        & {
            "ch_code",
            "account_number",
            "account_no",
            "relation",
            "sanitized_remark",
        }
    )


def _row_values(row: ProcessedRow) -> tuple[object, ...]:
    """Convert one processed row into the stable mapped export payload."""
    return (
        row.ch_code,
        row.account_number,
        row.relation.value,
        row.csu_rank,
        row.selected_rfd,
        row.sanitized_remark,
        row.remark_length,
        row.disposition.value,
    )


def _write_rows(
    workbook: Any,
    regions: dict[str, list[MappingRegion]],
    rows: list[ProcessedRow],
) -> None:
    """Write row values only inside validated mapping rectangles."""
    for sheet_name in TARGET_SHEETS:
        worksheet = workbook[sheet_name]
        remaining = list(rows)
        for region in regions[sheet_name]:
            if region.width < len(ROW_VALUE_FIELDS):
                raise RenderError(
                    f"Output_Mapping_Regions in {sheet_name} is too narrow: "
                    f"expected at least {len(ROW_VALUE_FIELDS)} columns, got {region.width}"
                )
            start_row = (
                region.min_row + 1
                if _looks_like_header(worksheet, region)
                else region.min_row
            )
            capacity = region.max_row - start_row + 1
            if capacity < 0:
                raise RenderError(f"Output_Mapping_Regions in {sheet_name} has invalid dimensions")
            chunk = remaining[:capacity]
            remaining = remaining[capacity:]
            for offset, row in enumerate(chunk):
                values = _row_values(row)
                target_row = start_row + offset
                for column_offset, value in enumerate(values):
                    worksheet.cell(
                        row=target_row,
                        column=region.min_col + column_offset,
                        value=value,
                    )
            if not remaining:
                break
        if remaining:
            raise RenderError(
                f"Output_Mapping_Regions in {sheet_name} cannot hold {len(rows)} rows"
            )


def _temporary_path(directory: Path, suffix: str = ".xlsx") -> Path:
    """Create a unique temporary path inside the output directory."""
    handle = tempfile.NamedTemporaryFile(
        mode="wb",
        suffix=suffix,
        prefix=".mc03-render-",
        dir=directory,
        delete=False,
    )
    path = Path(handle.name)
    handle.close()
    return path


def render_workbooks(result: RunResult, template_path: str, out_dir: str) -> dict[str, str]:
    """Render bank and internal workbooks from the fixed RCBC template.

    Only cells within discovered mapping regions are assigned.  The template is
    loaded with ``data_only=False`` so formula expressions remain expressions;
    openpyxl does not evaluate them.  Both temporary outputs must succeed
    before either final output path is replaced, leaving prior outputs intact on
    template/region/render failure.
    """
    template = Path(template_path)
    if not template.is_file():
        raise RenderError(f"fixed RCBC template is missing or unreadable: {template}")
    output_directory = Path(out_dir)
    output_directory.mkdir(parents=True, exist_ok=True)
    try:
        check_workbook = load_workbook(template, data_only=False)
        _discover_regions(check_workbook)
        check_workbook.close()
    except RenderError:
        raise
    except Exception as exc:
        raise RenderError(f"fixed RCBC template is missing or unreadable: {template}") from exc

    stem = _safe_output_stem(result.run_id)
    final_paths = {
        "bank_csr": output_directory / f"{stem}_bank_csr.xlsx",
        "internal_csr": output_directory / f"{stem}_internal_csr.xlsx",
    }
    temporary_paths: dict[str, Path] = {}
    try:
        for key in ("bank_csr", "internal_csr"):
            temporary = _temporary_path(output_directory)
            temporary_paths[key] = temporary
            workbook = load_workbook(template, data_only=False)
            regions = _discover_regions(workbook)
            # Use one shared row population for both fixed-template views.  The
            # two files are separate workbook instances, so no cross-file state
            # or formula evaluation can occur.
            _write_rows(workbook, regions, result.clean_rows)
            workbook.save(temporary)
            workbook.close()
        for key, temporary in temporary_paths.items():
            os.replace(temporary, final_paths[key])
    except RenderError:
        raise
    except Exception as exc:
        raise RenderError(f"workbook render failed: {type(exc).__name__}") from exc
    finally:
        for temporary in temporary_paths.values():
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                continue
    return {key: str(path) for key, path in final_paths.items()}


__all__ = [
    "MappingRegion",
    "OUTPUT_MAPPING_MARKER",
    "OUTPUT_MAPPING_REGIONS",
    "ROW_VALUE_FIELDS",
    "RenderError",
    "TARGET_SHEETS",
    "render_workbooks",
]
