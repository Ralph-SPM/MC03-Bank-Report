"""Focused confidence tests for the task-5 RCBC vertical slice."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pyzipper
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from mc03.domain.models import Disposition, ProcessedRow, Relation, RunResult
from mc03.persistence.run_store import RunStore
from mc03.services.archive import build_encrypted_archive, derive_archive_password
from mc03.services.orchestrator import process_run
from mc03.services.rendering import render_workbooks
from mc03.services.web import create_demo_app
from mc03.settings import RuntimeSettings


def _save_volare(path: Path, *, missing_rank: bool = False) -> None:
    """Create a minimal Volare workbook with a REF lookup sheet."""
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(["call date", "status", "remarks", "disposition"])
    data.append(
        [
            date(2026, 9, 2),
            "Completed",
            "TYPE OF RFD: Explicit | RFD: Medical | DETAILED RFD: illness | REMARKS: Follow up",
            "Known" if missing_rank else "Known",
        ]
    )
    reference = workbook.create_sheet("REF")
    reference.append(["disposition", "rank"])
    if not missing_rank:
        reference.append(["Known", 1])
    workbook.save(path)


def _save_field(path: Path) -> None:
    """Create a minimal filtered Field Result workbook."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "CH code",
            "bank",
            "status",
            "visit_date",
            "Contact Person",
            "Relation",
            "Statement",
            "client sentiment",
            "unit sentiment",
        ]
    )
    sheet.append(
        [
            "CH-1",
            "RCBC Auto Loan",
            "Completed",
            date(2026, 9, 2),
            "Jane Doe",
            "parent",
            "Promised to pay Friday",
            "positive",
            "positive",
        ]
    )
    workbook.save(path)


def _save_master(path: Path) -> None:
    """Create a valid single-match Master File workbook."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["CH code", "Account Number", "status"])
    sheet.append(["CH-1", "ACC-1", "RETAIN"])
    workbook.save(path)


def _save_template(path: Path) -> None:
    """Create a fixed-template-shaped workbook with mapping markers and metadata."""
    workbook = Workbook()
    first = workbook.active
    first.title = "FIELD RSULT"
    for sheet_name in ("FIELD RSULT", "DRR"):
        sheet = (
            workbook[sheet_name]
            if sheet_name in workbook.sheetnames
            else workbook.create_sheet(sheet_name)
        )
        sheet["A1"] = "Output_Mapping_Regions=A2:H10"
        for column, value in enumerate(
            [
                "CH code",
                "Account Number",
                "Relation",
                "CSU",
                "RFD",
                "Remark",
                "Length",
                "Disposition",
            ],
            start=1,
        ):
            sheet.cell(2, column, value)
        sheet["J1"] = "=1+1"
    metadata = workbook.create_sheet("REFERENCE")
    metadata.sheet_state = "hidden"
    metadata["A1"] = "=2+2"
    workbook.create_sheet("REF")
    workbook.create_sheet("MASTERLIST")
    workbook.save(path)


def test_run_store_overwrites_and_does_not_persist() -> None:
    """The store replaces by id and a new instance starts empty."""
    store = RunStore()
    first = RunResult(run_id="run-1")
    second = RunResult(run_id="run-1", excluded_count=2)
    store.store(first)
    assert store.get("run-1") is first
    store.store(second)
    assert store.retrieve("run-1") is second
    assert store.get("missing") is None
    assert len(RunStore()) == 0


def test_process_render_and_archive_round_trip(tmp_path: Path) -> None:
    """A real workbook set traverses the lazy pipeline and secure artifact path."""
    volare = tmp_path / "volare.xlsx"
    field = tmp_path / "field.xlsx"
    master = tmp_path / "master.xlsx"
    template = tmp_path / "template.xlsx"
    _save_volare(volare)
    _save_field(field)
    _save_master(master)
    _save_template(template)

    store = RunStore()
    result = process_run(
        date(2026, 9, 2),
        date(2026, 9, 3),
        volare,
        field,
        master,
        run_store=store,
    )
    assert result.attribution == "DA"
    assert len(result.clean_rows) == 1
    assert not result.review_rows
    assert len(result.clean_rows) + len(result.review_rows) + result.excluded_count == 1

    rendered = render_workbooks(result, str(template), str(tmp_path / "artifacts"))
    assert set(rendered) == {"bank_csr", "internal_csr"}
    output = load_workbook(rendered["bank_csr"], data_only=False)
    assert output["FIELD RSULT"]["A3"].value == "CH-1"
    assert output["FIELD RSULT"]["J1"].value == "=1+1"
    assert output["REFERENCE"].sheet_state == "hidden"

    archive = build_encrypted_archive(
        list(rendered.values()),
        str(tmp_path / "artifacts" / "package.zip"),
        result.date_from,
    )
    assert derive_archive_password(result.date_from) == "SEP2026"
    with pyzipper.AESZipFile(archive, "r") as zipped:
        zipped.setpassword(b"SEP2026")
        assert set(zipped.namelist()) == {Path(path).name for path in rendered.values()}
        assert zipped.read(Path(next(iter(rendered.values()))).name)


def test_review_decisions_are_isolated_and_export_gates(tmp_path: Path) -> None:
    """Only the targeted exception changes and unresolved rows gate export."""
    settings = RuntimeSettings(storage_root=tmp_path / "storage")
    app = create_demo_app(settings)
    result = RunResult(
        run_id="review-run",
        review_rows=[
            ProcessedRow(
                ch_code="CH-1",
                account_number="ACC-1",
                relation=Relation.UNKNOWN,
                csu_rank=0,
                selected_rfd=None,
                sanitized_remark="first",
                remark_length=5,
                disposition=Disposition.REVIEW_ONLY,
                source_row_ref="row 1",
            ),
            ProcessedRow(
                ch_code="CH-2",
                account_number="ACC-2",
                relation=Relation.UNKNOWN,
                csu_rank=0,
                selected_rfd=None,
                sanitized_remark="second",
                remark_length=6,
                disposition=Disposition.REVIEW_ONLY,
                source_row_ref="row 2",
            ),
        ],
    )
    app.state.run_store.store(result)
    client = TestClient(app, follow_redirects=False)
    page = client.get("/review/review-run")
    assert page.status_code == 200
    assert "first" in page.text and "second" in page.text
    gated = client.post("/export/review-run/generate")
    assert gated.status_code == 409
    edited = client.post(
        "/review/review-run/decision",
        data={"row_id": "0", "action": "edit", "csu_rank": "3", "rfd": "Medical"},
    )
    assert edited.status_code == 303
    assert result.review_rows[0].csu_rank == 3
    assert result.review_rows[0].selected_rfd == "Medical"
    assert result.review_rows[1].csu_rank == 0
    approved = client.post(
        "/review/review-run/decision",
        data={"row_id": "0", "action": "approve"},
    )
    assert approved.status_code == 303
    assert len(result.clean_rows) == 1
    assert len(result.review_rows) == 1
    invalid = client.post(
        "/review/review-run/decision",
        data={"row_id": "99", "action": "exclude"},
    )
    assert invalid.status_code == 400


def test_process_uploads_run_and_redirects_to_review(tmp_path: Path) -> None:
    """The successful upload route runs the pipeline and returns a 303 review redirect."""
    volare = tmp_path / "volare.xlsx"
    field = tmp_path / "field.xlsx"
    master = tmp_path / "master.xlsx"
    _save_volare(volare)
    _save_field(field)
    _save_master(master)
    app = create_demo_app(RuntimeSettings(storage_root=tmp_path / "storage"))
    client = TestClient(app, follow_redirects=False)
    xlsx_mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    response = client.post(
        "/process",
        data={"report_date_from": "2026-09-02", "report_date_to": "2026-09-03"},
        files={
            "volare_drr": ("volare.xlsx", volare.read_bytes(), xlsx_mime),
            "field_result": ("field.xlsx", field.read_bytes(), xlsx_mime),
            "master_file": ("master.xlsx", master.read_bytes(), xlsx_mime),
        },
    )
    assert response.status_code == 303
    location = response.headers["location"]
    assert location.startswith("/review/")
    run_id = location.rsplit("/", 1)[-1]
    stored = app.state.run_store.get(run_id)
    assert stored is not None
    assert stored.attribution == "DA"


def test_export_generation_exposes_only_three_completed_artifacts(tmp_path: Path) -> None:
    """Zero unresolved rows enable generation and download of the three artifacts."""
    storage_root = tmp_path / "storage"
    template_directory = storage_root / "templates"
    template_directory.mkdir(parents=True)
    template = template_directory / "RCBC AL_NEW-CSR_TEMPLATE JUNE 2026 (V1).xlsx"
    _save_template(template)
    volare = tmp_path / "volare.xlsx"
    field = tmp_path / "field.xlsx"
    master = tmp_path / "master.xlsx"
    _save_volare(volare)
    _save_field(field)
    _save_master(master)
    settings = RuntimeSettings(storage_root=storage_root, template_path=template_directory)
    app = create_demo_app(settings)
    result = process_run(
        date(2026, 9, 2),
        date(2026, 9, 3),
        volare,
        field,
        master,
        run_store=app.state.run_store,
    )
    assert app.state.run_store.get(result.run_id) is result
    client = TestClient(app, follow_redirects=False)
    response = client.post(f"/export/{result.run_id}/generate")
    assert response.status_code == 303, response.text
    export_page = client.get(f"/export/{result.run_id}")
    assert export_page.status_code == 200
    assert export_page.text.count("/download/") == 3
    for artifact in ("bank_csr", "internal_csr", "bank_package"):
        downloaded = client.get(f"/download/{result.run_id}/{artifact}")
        assert downloaded.status_code == 200
