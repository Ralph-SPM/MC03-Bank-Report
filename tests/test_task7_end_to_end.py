"""End-to-end dry-run coverage and deliverable verification for the RCBC Initial Demo portal.

Task 7.1 (end-to-end dry run with the historical sample file) and Task 7.2 (deliverable
verification assertions — clean-row length <= 200, Review_Only disposition invariants,
template formula preservation, AES-256 archive integrity) are executed and asserted here.
"""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import pyzipper
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from mc03.domain.models import Disposition
from mc03.services.archive import derive_archive_password
from mc03.services.web import create_demo_app
from mc03.settings import RuntimeSettings

HISTORICAL_SAMPLE_NAME = "Copy_of_DRR_TEMPLATE_UPDATED_FINAL_DRR_September_02-03_.xlsx"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _pipe_remark(statement: str) -> str:
    """Return a sanitized, structured remark suitable for the DRR upload."""
    return (
        "TYPE OF RFD: Explicit | RFD: Medical | DETAILED RFD: illness | "
        f"REMARKS: {statement}"
    )


def _save_sanitized_historical_sample(path: Path) -> None:
    """Create a small sanitized DRR-shaped sample covering clean, review, and noise rows."""
    workbook = Workbook()
    drr = workbook.active
    drr.title = "DRR"
    drr.append(["CH code", "call date", "status", "remarks", "disposition"])
    drr.append(
        [
            "CH-HIST-001",
            date(2026, 9, 2),
            "Completed",
            _pipe_remark("Contact confirmed a Friday payment promise."),
            "Known",
        ]
    )
    drr.append(
        [
            "CH-HIST-002",
            date(2026, 9, 3),
            "Completed",
            _pipe_remark("Contact requested a callback."),
            "Known",
        ]
    )
    drr.append(
        [
            "CH-HIST-003",
            date(2026, 9, 3),
            "New",
            _pipe_remark("Noise row retained only in the Field Result candidate set."),
            "Known",
        ]
    )
    reference = workbook.create_sheet("REF")
    reference.append(["disposition", "rank"])
    reference.append(["Known", 1])
    workbook.save(path)


def _save_field_result(path: Path) -> None:
    """Create the sanitized Field Result upload with three in-window candidates."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "CH code",
            "bank",
            "status",
            "visit_date",
            "remarks",
            "Contact Person",
            "Relation",
            "Statement",
            "client sentiment",
            "unit sentiment",
            "OB",
        ]
    )
    sheet.append(
        [
            "CH-HIST-001",
            "RCBC Auto Loan",
            "Completed",
            date(2026, 9, 2),
            _pipe_remark("Contact confirmed a Friday payment promise."),
            "Maria Santos",
            "parent",
            "Promised to pay on Friday",
            "positive",
            "positive",
            "field-value-1",
        ]
    )
    sheet.append(
        [
            "CH-HIST-002",
            "RCBC Auto Loan",
            "Completed",
            date(2026, 9, 3),
            _pipe_remark("Contact requested a callback."),
            "Neighbor Contact",
            "neighbor",
            "Will call back",
            "positive",
            "negative",
            "field-value-2",
        ]
    )
    sheet.append(
        [
            "CH-HIST-003",
            "RCBC Auto Loan",
            "New",
            date(2026, 9, 3),
            _pipe_remark("Noise row retained only in the Field Result candidate set."),
            "Noise Contact",
            "parent",
            "Noise disposition",
            "positive",
            "positive",
            "field-value-3",
        ]
    )
    workbook.save(path)


def _save_master_file(path: Path) -> None:
    """Create valid account matches for the clean/noise rows and omit the review row."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["CH code", "Account Number", "status"])
    sheet.append(["CH-HIST-001", "ACC-HIST-001", "RETAIN"])
    sheet.append(["CH-HIST-003", "ACC-HIST-003", "RESOLVED"])
    workbook.save(path)


def _save_fixed_template(path: Path) -> None:
    """Create the fixed-template-shaped workbook used by artifact generation."""
    workbook = Workbook()
    field_result = workbook.active
    field_result.title = "FIELD RSULT"
    for sheet_name in ("FIELD RSULT", "DRR"):
        sheet = (
            workbook[sheet_name]
            if sheet_name in workbook.sheetnames
            else workbook.create_sheet(sheet_name)
        )
        sheet["A1"] = "Output_Mapping_Regions=A2:H10"
        for column, value in enumerate(
            [
                "CH code",
                "Account Number",
                "Relation",
                "CSU",
                "RFD",
                "Remark",
                "Length",
                "Disposition",
            ],
            start=1,
        ):
            sheet.cell(2, column, value)
        sheet["J1"] = "=1+1"
    metadata = workbook.create_sheet("REFERENCE")
    metadata.sheet_state = "hidden"
    metadata["A1"] = "=2+2"
    workbook.create_sheet("REF")
    workbook.create_sheet("MASTERLIST")
    workbook.save(path)


def _assert_reconciled(result: object, candidate_count: int) -> None:
    """Assert that every candidate has exactly one terminal disposition bucket."""
    assert (
        len(result.clean_rows) + len(result.review_rows) + result.excluded_count
        == candidate_count
    )


def test_screen_1_ui_template_scope_trimming_and_dynamic_active_window(tmp_path: Path) -> None:
    """Screen 1 UI template must reflect approved scope trimming and dynamic active window."""
    app = create_demo_app(RuntimeSettings(storage_root=tmp_path / "storage"))
    client = TestClient(app, follow_redirects=False)

    response = client.get("/")
    assert response.status_code == 200
    html = response.text

    # 1. Local Mutex indicator removed
    assert "Local Mutex" not in html

    # 2. SHA-256 config snapshot line removed
    assert "SHA-256 Verified Config Snapshot" not in html

    # 3. Processing engine updated from WAL Streaming to In-Memory, Single Run
    assert "In-Memory WAL Streaming" not in html
    assert "In-Memory, Single Run" in html

    # 4. Active Window is dynamically bound to date inputs
    assert 'id="active-window-display"' in html
    assert 'id="active-window-text"' in html
    assert "2025-05-18 00:00:00" not in html


def test_historical_sample_upload_review_export_and_download(tmp_path: Path) -> None:
    """Drive a sanitized historical-shaped run through every portal stage."""
    historical_sample = tmp_path / HISTORICAL_SAMPLE_NAME
    field_result = tmp_path / "field_result.xlsx"
    master_file = tmp_path / "master_file.xlsx"
    storage_root = tmp_path / "storage"
    template_directory = storage_root / "templates"
    template = template_directory / "RCBC AL_NEW-CSR_TEMPLATE JUNE 2026 (V1).xlsx"

    _save_sanitized_historical_sample(historical_sample)
    _save_field_result(field_result)
    _save_master_file(master_file)
    template_directory.mkdir(parents=True)
    _save_fixed_template(template)

    app = create_demo_app(
        RuntimeSettings(storage_root=storage_root, template_path=template_directory)
    )
    client = TestClient(app, follow_redirects=False)
    response = client.post(
        "/process",
        data={"report_date_from": "2026-09-02", "report_date_to": "2026-09-03"},
        files={
            "volare_drr": (historical_sample.name, historical_sample.read_bytes(), XLSX_MIME),
            "field_result": (field_result.name, field_result.read_bytes(), XLSX_MIME),
            "master_file": (master_file.name, master_file.read_bytes(), XLSX_MIME),
        },
    )

    assert response.status_code == 303
    review_location = response.headers["location"]
    assert review_location.startswith("/review/")
    run_id = review_location.rsplit("/", 1)[-1]
    result = app.state.run_store.get(run_id)
    assert result is not None
    assert result.attribution == "DA"
    _assert_reconciled(result, candidate_count=3)
    assert len(result.clean_rows) == 1
    assert len(result.review_rows) == 1
    assert result.excluded_count == 1

    review_page = client.get(review_location)
    assert review_page.status_code == 200
    assert "CH-HIST-002" in review_page.text
    assert "CH-HIST-001" not in review_page.text
    assert "Unresolved Exceptions: <strong>1</strong>" in review_page.text

    edit_response = client.post(
        f"/review/{run_id}/decision",
        data={
            "row_id": "0",
            "action": "edit",
            "csu_rank": "CP+UP",
            "rfd": "Medical",
        },
    )
    assert edit_response.status_code == 303
    assert result.review_rows[0].csu_rank == 3
    assert result.review_rows[0].selected_rfd == "Medical"

    approve_response = client.post(
        f"/review/{run_id}/decision",
        data={"row_id": "0", "action": "approve"},
    )
    assert approve_response.status_code == 303
    assert not result.review_rows
    assert all(row.disposition is Disposition.CLEAN for row in result.clean_rows)
    _assert_reconciled(result, candidate_count=3)

    export_page = client.get(f"/export/{run_id}")
    assert export_page.status_code == 200
    assert "Included Rows" in export_page.text
    assert "Excluded Noise" in export_page.text
    assert "Unresolved Exceptions" in export_page.text
    assert "Unresolved Exceptions</dt><dd>0</dd>" in export_page.text

    generation_response = client.post(f"/export/{run_id}/generate")
    assert generation_response.status_code == 303

    generated_page = client.get(f"/export/{run_id}")
    assert generated_page.status_code == 200
    expected_artifacts = {"bank_csr", "internal_csr", "bank_package"}
    assert all(
        f"/download/{run_id}/{artifact}" in generated_page.text
        for artifact in expected_artifacts
    )
    assert generated_page.text.count(f"/download/{run_id}/") == 3
    assert set(result.artifacts) == expected_artifacts

    for artifact in expected_artifacts:
        download = client.get(f"/download/{run_id}/{artifact}")
        assert download.status_code == 200
        assert download.content


def test_task7_2_deliverable_verification_assertions(tmp_path: Path) -> None:
    """Task 7.2 deliverable verification assertions against historical sample run.

    Verifies the four required properties:
    1. Clean-row remark length <= 200
    2. Review_Only disposition invariants (state, gating, reconciliation)
    3. Template formula preservation (=1+1, hidden reference sheet state)
    4. AES-256 archive integrity (pyzipper password derive and decryption test)
    """
    historical_sample = tmp_path / HISTORICAL_SAMPLE_NAME
    field_result = tmp_path / "field_result.xlsx"
    master_file = tmp_path / "master_file.xlsx"
    storage_root = tmp_path / "storage"
    template_directory = storage_root / "templates"
    template = template_directory / "RCBC AL_NEW-CSR_TEMPLATE JUNE 2026 (V1).xlsx"

    _save_sanitized_historical_sample(historical_sample)
    _save_field_result(field_result)
    _save_master_file(master_file)
    template_directory.mkdir(parents=True)
    _save_fixed_template(template)

    app = create_demo_app(
        RuntimeSettings(storage_root=storage_root, template_path=template_directory)
    )
    client = TestClient(app, follow_redirects=False)

    # 1. Process historical sample run
    process_response = client.post(
        "/process",
        data={"report_date_from": "2026-09-02", "report_date_to": "2026-09-03"},
        files={
            "volare_drr": (historical_sample.name, historical_sample.read_bytes(), XLSX_MIME),
            "field_result": (field_result.name, field_result.read_bytes(), XLSX_MIME),
            "master_file": (master_file.name, master_file.read_bytes(), XLSX_MIME),
        },
    )
    assert process_response.status_code == 303
    run_id = process_response.headers["location"].rsplit("/", 1)[-1]
    result = app.state.run_store.get(run_id)
    assert result is not None

    # --- Property 2 Verification: Review_Only disposition invariants (Part A) ---
    _assert_reconciled(result, candidate_count=3)
    assert len(result.clean_rows) == 1
    assert len(result.review_rows) == 1
    assert result.excluded_count == 1

    # Invariant: Review_Only rows must carry REVIEW_ONLY disposition and non-empty exception_kinds
    for review_row in result.review_rows:
        assert review_row.disposition is Disposition.REVIEW_ONLY
        assert len(review_row.exception_kinds) > 0

    # Invariant: Unresolved review rows MUST gate artifact generation (returns 409)
    gated_response = client.post(f"/export/{run_id}/generate")
    assert gated_response.status_code == 409
    assert gated_response.text == "Unresolved exceptions must be zero before generation."

    # Resolve exception by editing and approving
    client.post(
        f"/review/{run_id}/decision",
        data={"row_id": "0", "action": "edit", "csu_rank": "CP+UP", "rfd": "Medical"},
    )
    approve_response = client.post(
        f"/review/{run_id}/decision",
        data={"row_id": "0", "action": "approve"},
    )
    assert approve_response.status_code == 303

    # --- Property 2 Verification: Review_Only disposition invariants (Part B) ---
    assert len(result.review_rows) == 0
    assert len(result.clean_rows) == 2
    assert all(row.disposition is Disposition.CLEAN for row in result.clean_rows)
    _assert_reconciled(result, candidate_count=3)

    # --- Property 1 Verification: Clean-row length <= 200 ---
    for clean_row in result.clean_rows:
        assert len(clean_row.sanitized_remark) <= 200
        assert clean_row.remark_length <= 200

    # Generate artifacts now that unresolved count is zero
    gen_response = client.post(f"/export/{run_id}/generate")
    assert gen_response.status_code == 303

    # Fetch downloads
    bank_csr_res = client.get(f"/download/{run_id}/bank_csr")
    internal_csr_res = client.get(f"/download/{run_id}/internal_csr")
    package_res = client.get(f"/download/{run_id}/bank_package")

    assert bank_csr_res.status_code == 200
    assert internal_csr_res.status_code == 200
    assert package_res.status_code == 200

    # Property 1 cell check on output workbooks
    for wb_bytes in (bank_csr_res.content, internal_csr_res.content):
        wb = load_workbook(io.BytesIO(wb_bytes), data_only=True)
        for sheet in (wb["FIELD RSULT"], wb["DRR"]):
            # Remark is column 6 in header row 2
            for row in range(3, sheet.max_row + 1):
                val = sheet.cell(row, 6).value
                if val is not None:
                    assert len(str(val)) <= 200

    # --- Property 3 Verification: Template formula preservation ---
    for wb_bytes in (bank_csr_res.content, internal_csr_res.content):
        wb = load_workbook(io.BytesIO(wb_bytes), data_only=False)
        # Verify formula cell in data sheets preserved starting with '='
        assert str(wb["FIELD RSULT"]["J1"].value).startswith("=")
        assert wb["FIELD RSULT"]["J1"].value == "=1+1"
        assert wb["DRR"]["J1"].value == "=1+1"
        # Verify hidden reference sheet & formula state preserved
        assert wb["REFERENCE"].sheet_state == "hidden"
        assert wb["REFERENCE"]["A1"].value == "=2+2"

    # --- Property 4 Verification: AES-256 archive integrity ---
    package_bytes = package_res.content
    assert len(package_bytes) > 0
    expected_password = derive_archive_password(result.date_from)
    assert expected_password == "SEP2026"

    with pyzipper.AESZipFile(io.BytesIO(package_bytes), "r") as zipped:
        zipped.setpassword(expected_password.encode("utf-8"))
        # testzip() returns None if all zip CRC / AES decryptions pass
        assert zipped.testzip() is None
        file_names = zipped.namelist()
        assert len(file_names) == 2
        for name in file_names:
            extracted_data = zipped.read(name)
            assert len(extracted_data) > 0
            # Ensure extracted workbooks open as valid excel workbooks
            test_wb = load_workbook(io.BytesIO(extracted_data), read_only=True)
            test_wb.close()
