"""Unit tests for Remarks Intelligence Lab & Benchmark View."""

from __future__ import annotations

import io
import pandas as pd
import pytest
from starlette.testclient import TestClient

from mc03.services.remarks_lab.parser import parse_field_result_sheet, RawRemarkRow
from mc03.services.remarks_lab.cleaner import clean_remark
from mc03.services.remarks_lab.classifier import classify_contact
from mc03.services.remarks_lab.ranker import rank_rfd_and_csu
from mc03.services.remarks_lab.trimmer import trim_and_format
from mc03.services.remarks_lab.pipeline import RemarksLabPipeline
from mc03.services.web import create_demo_app


def test_cleaner_strips_prohibited_tags_and_phones():
    text = "Spoke to neighbor BCAL L3 OBD 09171234567 visit result ch moved to prov"
    res = clean_remark(text)
    assert res.is_modified
    assert "BCAL" in res.detected_prohibited_tags
    assert "L3" in res.detected_prohibited_tags
    assert "OBD" in res.detected_prohibited_tags
    assert "PHONE" in res.detected_prohibited_tags
    assert "09171234567" not in res.cleaned
    assert "BCAL" not in res.cleaned


def test_classifier_extended_filipino_and_guardrail():
    # 1. Representative Tagalog
    rep1 = classify_contact("Maria Santos", "kapatid")
    assert rep1.category_label == "Representative"
    assert rep1.normalized_role.lower() == "kapatid"

    rep2 = classify_contact("Juan Dela Cruz", "nanay")
    assert rep2.category_label == "Representative"

    # 2. Informant Tagalog
    inf1 = classify_contact("Kuya Guard", "sekyu")
    assert inf1.category_label == "Informant"
    assert inf1.guardrail_applied

    inf2 = classify_contact("Pedro", "kapitbahay")
    assert inf2.category_label == "Informant"
    assert inf2.guardrail_applied

    # 3. Cardholder / Borrower
    ch = classify_contact("Self", "Borrower")
    assert ch.category_label in ("Borrower", "Cardholder")


def test_ranker_guardrail_informant_never_representative_refused():
    rank_rep = rank_rfd_and_csu("borrower refused to talk", "borrower refused to talk", "Representative")
    assert rank_rep.rfd_code == "REPRESENTATIVE REFUSED TO DISCLOSE RFD"

    rank_inf = rank_rfd_and_csu("neighbor said ch moved out", "neighbor said ch moved out", "Informant")
    assert rank_inf.rfd_code != "REPRESENTATIVE REFUSED TO DISCLOSE RFD"
    assert rank_inf.rfd_code == "Moved Out"


def test_trimmer_preamble_and_length():
    long_statement = "FIELD VISIT RESULT: " + ("The borrower promises to settle balance next week upon release of harvest funds. " * 5)
    trimmed = trim_and_format(long_statement)
    assert trimmed.preamble_stripped
    assert trimmed.trimmed_length <= 200
    assert trimmed.fits_within_200


def test_remarks_lab_pipeline_synthetic_excel():
    # Create synthetic DataFrame
    data = {
        "Account Number": ["1001", "1002", "1003"],
        "CH Code": ["CH01", "CH02", "CH03"],
        "Contact Person": ["Juana Santos", "Brgy Tanod Pedro", "Cardholder"],
        "Contact Relation": ["Asawa", "Barangay Tanod", "Self"],
        "Remarks": [
            "BCAL spoke with asawa, PTP on Friday 09181234567",
            "OBD CH moved to Cavite per tanod",
            "CH refused to pay"
        ],
        "CSU": ["Representative PTP", "Moved Out / Relocated", "Borrower Refused"],
        "RFD": ["PTP", "Moved", "Refused to pay"]
    }
    df = pd.DataFrame(data)
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="FIELD RSULT", index=False)
    buffer.seek(0)

    pipeline = RemarksLabPipeline()
    report = pipeline.process_file(buffer)

    assert report.total_rows == 3
    assert report.sanitized_tag_count >= 2
    assert report.representative_count == 1
    assert report.informant_count == 1
    assert report.cardholder_count == 1
    assert len(report.rows) == 3


def test_remarks_lab_endpoints():
    app = create_demo_app()
    client = TestClient(app)

    # GET /remarks-lab
    get_res = client.get("/remarks-lab")
    assert get_res.status_code == 200
    assert "Remark Intelligence Lab" in get_res.text
    assert "Upload FIELD RSULT Workbook" in get_res.text

    # POST /remarks-lab/analyze with synthetic workbook
    df = pd.DataFrame({
        "Account Number": ["9999"],
        "Contact Person": ["Maria"],
        "Contact Relation": ["Ate"],
        "Remarks": ["L3 PTP next Tuesday"]
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="FIELD RSULT", index=False)
    data = buf.getvalue()

    post_res = client.post(
        "/remarks-lab/analyze",
        files={"workbook": ("test.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert post_res.status_code == 200
    assert "Rows Analyzed" in post_res.text
    assert "9999" in post_res.text
    assert "Representative" in post_res.text


def test_remarks_lab_template_shaped_workbook_parsing():
    """Verify parsing workbooks with metadata headers like RCBC AL_NEW-CSR_TEMPLATE."""
    buf = io.BytesIO()
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "FIELD RSULT"
    ws["A1"] = "Output_Mapping_Regions=A2:H10"
    headers = ["CH code", "Account Number", "Relation", "CSU", "RFD", "Remark", "Length", "Disposition"]
    for col, val in enumerate(headers, start=1):
        ws.cell(2, col, val)

    ws.append(["CH-999", "ACC-999", "parent", "3", "Medical", "This is a long remark over 200 chars. " * 8, 280, "Clean"])
    wb.save(buf)
    data = buf.getvalue()

    pipeline = RemarksLabPipeline()
    report = pipeline.process_file(io.BytesIO(data))
    assert report.total_rows == 1
    assert report.char_overflow_prevented_count == 1
    assert report.rows[0].account_number == "ACC-999"
    assert report.rows[0].truncated is True


def test_remarks_lab_date_filtering():
    from datetime import date
    df = pd.DataFrame({
        "Visit Date": ["2026-08-15", "2026-09-02", "2026-09-03", "2026-09-10"],
        "Account Number": ["ACC-1", "ACC-2", "ACC-3", "ACC-4"],
        "Contact Person": ["Maria", "Pedro", "Juan", "Ana"],
        "Contact Relation": ["Ate", "Kuya", "Tatay", "Nanay"],
        "Remarks": [
            "PTP on Aug 20",
            "PTP on Sep 5",
            "Left notice with neighbor",
            "CH not at residence"
        ]
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="FIELD RSULT", index=False)
    data = buf.getvalue()

    # 1. Pipeline direct test
    pipeline = RemarksLabPipeline()
    report = pipeline.process_file(
        io.BytesIO(data),
        date_from=date(2026, 9, 2),
        date_to=date(2026, 9, 3)
    )
    assert report.total_rows == 2
    assert report.total_unfiltered_rows == 4
    assert [r.account_number for r in report.rows] == ["ACC-2", "ACC-3"]
    assert report.rows[0].row_date == "2026-09-02"
    assert report.rows[1].row_date == "2026-09-03"

    # 2. HTTP POST endpoint test with date_from and date_to
    app = create_demo_app()
    client = TestClient(app)
    post_res = client.post(
        "/remarks-lab/analyze",
        data={"date_from": "2026-09-02", "date_to": "2026-09-03"},
        files={"workbook": ("csr_template.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert post_res.status_code == 200
    assert "Date Window Filter" in post_res.text
    assert "2026-09-02" in post_res.text
    assert "2026-09-03" in post_res.text
    assert "isolated from" in post_res.text
    assert "ACC-2" in post_res.text
    assert "ACC-3" in post_res.text
    assert "ACC-1" not in post_res.text
    assert "ACC-4" not in post_res.text


def test_rcbc_csu_taxonomy_and_rfd_hierarchy():
    from engine.classifier import classify_csu, classify_rfd, rank_rfd_and_csu

    # 1. Repossession -> CLIENT POSITIVE/UNIT POSITIVE (WITH Actual Contact - Both)
    repo_res = rank_rfd_and_csu("Unit successfully repossessed Sept 1, 2026", "Unit successfully repossessed Sept 1, 2026", "Borrower")
    assert repo_res.csu_status == "CLIENT POSITIVE/UNIT POSITIVE (WITH Actual Contact - Both)"
    assert classify_csu("Voluntary surrender of mortgaged unit", "Borrower") == "CLIENT POSITIVE/UNIT POSITIVE (WITH Actual Contact - Both)"

    # 2. Payment Promise -> WITH Actual Contact - Both and BORROWER REFUSED TO DISCLOSE RFD
    ptp_res = rank_rfd_and_csu("Client PTP on Sept 4 for 15,000", "Client PTP on Sept 4 for 15,000", "Cardholder")
    assert ptp_res.csu_status == "CLIENT POSITIVE/UNIT POSITIVE (WITH Actual Contact - Both)"
    assert ptp_res.rfd_code == "BORROWER REFUSED TO DISCLOSE RFD"

    # 3. Account Resolved -> CLIENT POSITIVE/UNIT POSITIVE (WITHOUT Actual Contact - Client)
    res_res = rank_rfd_and_csu("Neg; per agent, account already resolved", "Neg; per agent, account already resolved", "none")
    assert res_res.csu_status == "CLIENT POSITIVE/UNIT POSITIVE (WITHOUT Actual Contact - Client)"

    # 4. No Contact / House Closed (No Informant) -> CLIENT NEGATIVE/UNIT NEGATIVE and empty RFD
    nc_res = rank_rfd_and_csu("House closed gate padlocked nobody home", "House closed gate padlocked nobody home", "Unknown")
    assert nc_res.csu_status == "CLIENT NEGATIVE/UNIT NEGATIVE (FOR FURTHER VISIT/PROBING)"
    assert nc_res.rfd_code == ""

    # 5. Explicit Hardships
    med_res = rank_rfd_and_csu("Borrower admitted in hospital due to medical surgery", "Borrower admitted in hospital due to medical surgery", "Cardholder")
    assert med_res.rfd_code == "Medical Expense"

    sal_res = rank_rfd_and_csu("Delayed salary from employer, promised to coordinate next week", "Delayed salary from employer, promised to coordinate next week", "Cardholder")
    assert sal_res.rfd_code == "Delayed Salary"

    bus_res = rank_rfd_and_csu("Business slowdown in store, low sales this quarter", "Business slowdown in store, low sales this quarter", "Cardholder")
    assert bus_res.rfd_code == "Business Slowdown"

    third_res = rank_rfd_and_csu("Car was transferred to third-party user under pasalo arrangement", "Car was transferred to third-party user under pasalo arrangement", "Representative")
    assert third_res.rfd_code == "Third-Party User"

    # 6. Representative contact with no hardship
    rep_res = rank_rfd_and_csu("Spoke with mother, refused to provide update", "Spoke with mother, refused to provide update", "Representative")
    assert rep_res.rfd_code == "REPRESENTATIVE REFUSED TO DISCLOSE RFD"

    # 7. Informant with moved out
    inf_res = rank_rfd_and_csu("Neighbor said cardholder moved out to Cavite", "Neighbor said cardholder moved out to Cavite", "Informant")
    assert inf_res.rfd_code == "Moved Out"

    # Never output FIELD VISIT or PTP as RFD
    for rfd in [repo_res.rfd_code, ptp_res.rfd_code, rep_res.rfd_code, inf_res.rfd_code]:
        assert rfd not in ["FIELD VISIT", "PTP", ""]


def test_agency_preamble_removal_and_deduplication():
    from engine.trimmer import trim_and_format, strip_preamble, deduplicate_phrases

    # 1. Internal preamble removal
    sample_with_header = "FIELD VISIT RESULT: Unit successfully repossessed Sept 1, 2026"
    stripped, was_stripped = strip_preamble(sample_with_header)
    assert was_stripped is True
    assert "FIELD VISIT RESULT" not in stripped
    assert "Unit successfully repossessed" in stripped

    # 2. Phrase deduplication
    repeated_text = "Neg; per agent, account already resolved. Neg; per agent, account already resolved. Neg; per agent, account already resolved."
    deduped = deduplicate_phrases(repeated_text)
    assert deduped.count("Neg; per agent, account already resolved") == 1

    # 3. Combined trimmer test (ECA header preserved)
    combined = "ECA AUTO_S.P. MADRID_HOME_09/02/2026 - Neg; per agent, account already resolved. Neg; per agent, account already resolved."
    res = trim_and_format(combined)
    assert res.trimmed_length <= 200
    assert "ECA AUTO_S.P. MADRID_HOME_09/02/2026" in res.trimmed_text
    assert res.trimmed_text.count("account already resolved") == 1


def test_rcbc_226_rows_benchmark_accuracy():
    """Verify CSU match > 80% and accurate RFD mapping on representative 226 records."""
    rows_data = []
    # 50 PTP rows
    for i in range(50):
        rows_data.append({
            "Account Number": f"ACC-{i}",
            "Relation": "Borrower",
            "Remarks": f"ECA AUTO_S.P. MADRID_HOME_09/02/2026 - Spoke with client, PTP on Sept {i%20 + 1}",
            "CSU": "PTP",
            "RFD": "BORROWER REFUSED TO DISCLOSE RFD",
        })
    # 30 REPO rows
    for i in range(50, 80):
        rows_data.append({
            "Account Number": f"ACC-{i}",
            "Relation": "Self",
            "Remarks": "Unit successfully repossessed Sept 1, 2026 by recovery team",
            "CSU": "REPO",
            "RFD": "BORROWER REFUSED TO DISCLOSE RFD",
        })
    # 20 RESOLVED rows
    for i in range(80, 100):
        rows_data.append({
            "Account Number": f"ACC-{i}",
            "Relation": "",
            "Remarks": "Neg; per agent, account already resolved. Neg; per agent, account already resolved.",
            "CSU": "RESOLVED",
            "RFD": "NO CLIENT / REPRESENTATIVE REACHED",
        })
    # 100 House closed / No contact rows
    for i in range(100, 200):
        rows_data.append({
            "Account Number": f"ACC-{i}",
            "Relation": "Neighbor",
            "Remarks": "House closed gate locked, neighbor said not around",
            "CSU": "CLIENT NEGATIVE/UNIT NEGATIVE (FOR FURTHER VISIT/PROBING)",
            "RFD": "NO CLIENT / REPRESENTATIVE REACHED",
        })
    # 26 Explicit hardship rows
    hardships = [
        ("medical surgery in hospital", "Medical Expense"),
        ("delayed salary from employer", "Delayed Salary"),
        ("business slowdown low revenue", "Business Slowdown"),
        ("diversion of funds emergency", "Diversion of Funds"),
        ("delayed collection from contractor", "Delayed Collection"),
        ("delayed remittance from OFW brother", "Delayed Remittance"),
        ("third-party user pasalo arrangement", "Third-Party User"),
    ]
    for i in range(200, 226):
        phrase, exp_rfd = hardships[i % len(hardships)]
        rows_data.append({
            "Account Number": f"ACC-{i}",
            "Relation": "Client",
            "Remarks": f"Borrower explained default due to {phrase}, PTP next month",
            "CSU": "PTP",
            "RFD": exp_rfd,
        })

    df = pd.DataFrame(rows_data)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="FIELD RSULT", index=False)
    buf.seek(0)

    pipeline = RemarksLabPipeline()
    report = pipeline.process_file(buf)

    assert report.total_rows == 226
    assert report.csu_accuracy_pct > 80.0
    assert report.rfd_accuracy_pct > 80.0
    # Specific row checks
    repo_row = next(r for r in report.rows if "repossessed" in r.raw_remarks)
    assert "UNIT POSITIVE" in repo_row.predicted_csu
    assert repo_row.csu_match is True

    ptp_row = next(r for r in report.rows if "PTP on Sept" in r.raw_remarks)
    assert "CLIENT POSITIVE" in ptp_row.predicted_csu
    assert ptp_row.predicted_rfd == "BORROWER REFUSED TO DISCLOSE RFD"
    assert ptp_row.csu_match is True
    assert ptp_row.rfd_match is True


def test_bug1_3tier_relation_classifier_and_eca():
    # 1. Borrower (Self / direct conversation)
    b1 = classify_contact("", "", "Client PTP to settle account on Friday")
    assert b1.category_label == "Borrower"
    assert b1.normalized_role == "Borrower (Self)"

    b2 = classify_contact("", "", "Spoke with borrower directly at residence")
    assert b2.category_label == "Borrower"

    # 2. Representative (In-law, blood relative)
    r1 = classify_contact("", "", "Spoke with client's sister-in-law regarding car loan")
    assert r1.category_label == "Representative"

    r2 = classify_contact("", "", "Per tita, client is out of the country")
    assert r2.category_label == "Representative"

    # 3. Informant (Guard, neighbor, tanod, colleague)
    i1 = classify_contact("", "", "Per building guard, client moved out last year")
    assert i1.category_label == "Informant"
    assert i1.guardrail_applied is True

    i2 = classify_contact("Kuya Guard", "sekyu", "Unit parked at garage; spoke with guard")
    assert i2.category_label == "Informant"
    assert i2.category_label != "Representative"

    # When remark doesn't say who was contacted, must be 'none' despite metadata
    res_none = classify_contact("Kuya Guard", "sekyu", "Unit parked at garage")
    assert res_none.category_label == "none"
    assert res_none.normalized_role == "none"

    # 4. None reached (Never "Unknown")
    n1 = classify_contact("", "", "Gate closed, nobody answered intercom")
    assert n1.category_label == "none"
    assert n1.normalized_role == "none"

    # 5. Remark is sole source of truth: if remark doesn't mention mother, do not infer from metadata
    eca_rep = classify_contact(
        "", "",
        "ECA AUTO_S.P. MADRID_HOME_09/02/2026 Per agent, account already resolved.",
        extra_fields={"3RD PARTY LIST": "mother"}
    )
    assert eca_rep.category_label == "none"

    # But if remark DOES mention mother, classify as Representative
    eca_with_mother = classify_contact(
        "", "",
        "ECA AUTO_S.P. MADRID_HOME_09/02/2026 Spoke with mother, per agent account resolved.",
        extra_fields={"3RD PARTY LIST": "mother"}
    )
    assert eca_with_mother.category_label == "Representative"
    assert eca_with_mother.normalized_role == "Mother"


def test_bug2_detailed_rfd_3clause_and_hierarchies():
    from engine.classifier import rank_rfd_and_csu, classify_rfd, classify_csu

    # 1. DETAILED RFD 3-Clause Structure: [RFD clause]; [Contact clause]; [Core statement/promise]
    res_b = rank_rfd_and_csu("Client promised to pay 10k on Oct 15", "Client promised to pay 10k on Oct 15", "Borrower", "Borrower (Self)")
    parts_b = [p.strip() for p in res_b.detailed_rfd.split(";")]
    assert len(parts_b) == 3
    assert parts_b[0] == "BORROWER REFUSED TO DISCLOSE RFD"
    assert parts_b[1] == "Contact made with borrower directly"
    assert "Promised to pay" in parts_b[2]

    res_rep = rank_rfd_and_csu("Spoke with wife, account already resolved", "Spoke with wife, account already resolved", "Representative", "Wife")
    parts_rep = [p.strip() for p in res_rep.detailed_rfd.split(";")]
    assert len(parts_rep) == 3
    assert parts_rep[0] == "REPRESENTATIVE REFUSED TO DISCLOSE RFD"
    assert parts_rep[1] == "Contact made with representative (Wife)"
    assert "Account already resolved" in parts_rep[2]

    res_none = rank_rfd_and_csu("Gate closed, unlocated", "Gate closed, unlocated", "none reached", "none reached")
    parts_none = [p.strip() for p in res_none.detailed_rfd.split(";")]
    assert len(parts_none) == 3
    assert parts_none[0] == "NO CLIENT / REPRESENTATIVE REACHED"
    assert parts_none[1] == "No contact made"

    # 2. RFD Value Priority: Explicit hardship > Borrower refused > Representative refused > No client > Moved out
    assert classify_rfd("Client hospital surgery due to medical condition", "Borrower") == "Medical Expense"
    # Secondary order: medical expense > diversion of funds
    assert classify_rfd("Client had medical expense and diversion of funds", "Borrower") == "Medical Expense"

    # 3. CSU Hierarchy: Repo / Resolved > Client Positive + Unit Positive > Client Positive + Unit Negative > Client Negative
    assert classify_csu("Client surrendered unit to repo agent", "Borrower") == "CLIENT POSITIVE/UNIT POSITIVE (WITH Actual Contact - Both)"
    assert classify_csu("Address verified per neighbor, unit seen in garage", "Informant") == "CLIENT POSITIVE/UNIT POSITIVE (WITHOUT Actual Contact - Client)"
    assert classify_csu("Address verified per neighbor, unit not seen", "Informant") == "CLIENT POSITIVE/UNIT NEGATIVE (WITHOUT Actual Contact)"
    assert classify_csu("Client unknown in area, cannot locate", "none reached") == "CLIENT NEGATIVE/UNIT NEGATIVE (FOR FURTHER VISIT/PROBING)"


def test_bug3_llm_trimming_and_safety_fallback():
    from engine.trimmer import trim_and_format, fallback_clause_truncate

    # 1. Fallback clause truncation: preserves substantive findings without destructive cuts
    long_remark = (
        "Positive address verified by neighbor Mr. Santos; borrower is temporarily out of town for business trip until Monday; "
        "unit is parked inside covered garage with plate ABC-1234; neighbor advised agent to return Tuesday morning."
    )
    trimmed = fallback_clause_truncate(long_remark, max_chars=200, preferred_chars=181)
    assert len(trimmed) <= 181
    assert not trimmed.endswith(";")
    assert not trimmed.endswith(",")
    # Substantive content retained, NOT 7-41 chars
    assert len(trimmed) > 100
    assert "Positive address verified" in trimmed

    # 2. LLM Trimmer mock with retry
    call_count = 0

    def mock_llm_provider(text: str, is_retry: bool) -> str:
        nonlocal call_count
        call_count += 1
        if not is_retry:
            # First attempt returns > 200 chars
            return "Contacted borrower directly: discussed overdue payment terms and unit condition in extensive detail with multiple payment promises scheduled for next month across several installments exceeding two hundred characters."
        # Retry returns <= 181 chars
        return "Spoke with borrower directly; agreed to settle overdue payment next month."

    res = trim_and_format(long_remark, llm_provider=mock_llm_provider)
    assert call_count == 2  # Proves retry was invoked
    assert res.trimmed_length <= 181
    assert "Spoke with borrower directly" in res.trimmed_text


def test_export_to_excel_and_endpoint():
    pipeline = RemarksLabPipeline()
    df = pd.DataFrame({
        "Account Number": ["ACC-001", "ACC-002"],
        "CH Code": ["CH1", "CH2"],
        "Contact Person": ["Maria", "Pedro"],
        "Contact Relation": ["Asawa", "Neighbor"],
        "Remarks": [
            "BKAL Spoke with wife, PTP on Oct 5 09181234567",
            "Per neighbor, client moved out years ago"
        ],
        "COLLECTION STATUS UPDATE\nAGENT/ADMIN": [
            "CLIENT POSITIVE/UNIT NEGATIVE (WITH Actual Contact)",
            "CLIENT NEGATIVE/UNIT NEGATIVE (FOR FURTHER VISIT/PROBING)"
        ],
        "RFD": [
            "REPRESENTATIVE REFUSED TO DISCLOSE RFD",
            "MOVED OUT"
        ]
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="FIELD RSULT", index=False)
    data = buf.getvalue()

    report = pipeline.process_file(io.BytesIO(data))
    excel_out = pipeline.export_to_excel(report)
    assert excel_out is not None

    # Verify generated Excel content
    exported_df = pd.read_excel(excel_out, sheet_name="FIELD RESULT (EVALUATED)")
    assert len(exported_df) == 2
    assert "char checker" in exported_df.columns
    assert list(exported_df["char checker"].fillna("")) == ["", ""]
    assert "COLLECTION STATUS UPDATE" in exported_df.columns
    assert "RFD" in exported_df.columns

    # Test HTTP POST /remarks-lab/export endpoint
    app = create_demo_app()
    client = TestClient(app)
    resp = client.post(
        "/remarks-lab/export",
        files={"workbook": ("test.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "RCBC_FIELD_EVALUATED" in resp.headers["content-disposition"]


def test_export_to_excel_test_mode():
    pipeline = RemarksLabPipeline()
    df = pd.DataFrame({
        "Account Number": ["ACC-999"],
        "CH Code": ["CH9"],
        "Contact Person": ["Juan Dela Cruz"],
        "Contact Relation": ["Self"],
        "MESSAGE": ["Spoke with borrower; PTP next Monday"],
        "FINAL REMARKS": ["ECA AUTO_S.P. MADRID_HOME_09/02/2026 Spoke with borrower"],
        "COLLECTION STATUS UPDATE\nAGENT/ADMIN": ["CLIENT POSITIVE/UNIT POSITIVE (WITH Actual Contact)"],
        "RFD": ["BORROWER REFUSED TO DISCLOSE RFD"]
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="FIELD RSULT", index=False)
    data = buf.getvalue()

    report = pipeline.process_file(io.BytesIO(data), test_mode=True)
    excel_out = pipeline.export_to_excel(report)
    assert excel_out is not None

    exported_df = pd.read_excel(excel_out, sheet_name="FIELD RESULT (EVALUATED)")
    assert len(exported_df) == 1
    assert "Raw Remark" in exported_df.columns
    assert "Original CSU" in exported_df.columns
    assert "Original RFD" in exported_df.columns
    assert "CSU Why" in exported_df.columns
    assert "RFD Why" in exported_df.columns
    assert exported_df["Raw Remark"].iloc[0] == "Spoke with borrower; PTP next Monday"
    assert exported_df["Original CSU"].iloc[0] == "CLIENT POSITIVE/UNIT POSITIVE (WITH Actual Contact)"
    assert exported_df["Original RFD"].iloc[0] == "BORROWER REFUSED TO DISCLOSE RFD"
    assert exported_df["CSU Why"].iloc[0] != ""
    assert exported_df["RFD Why"].iloc[0] != ""

    # Test round-trip: re-importing the exported test file parses the whys
    excel_out.seek(0)
    reimported_report = pipeline.process_file(excel_out, test_mode=True, run_ai=False)
    assert len(reimported_report.rows) == 1
    assert reimported_report.rows[0].csu_reasoning == exported_df["CSU Why"].iloc[0]
    assert reimported_report.rows[0].rfd_reasoning == exported_df["RFD Why"].iloc[0]


def test_expanded_hardships_and_informant_guardrails():
    from engine.classifier import classify_contact, classify_rfd, classify_csu

    # 1. Calamity
    assert classify_rfd("Per client, residence was flooded due to typhoon") == "Calamity"

    # 2. Scammed
    assert classify_rfd("Client stated was scammed by coworker who took vehicle") == "Scammed"

    # 3. LTO Apprehension
    assert classify_rfd("Unit impounded by LTO; client coordinating for release") == "LTO Apprehension/No ORCR/HPG"

    # 4. Work Relocation
    assert classify_rfd("Per sibling, client no longer resides here, working in Palawan") == "Work Relocation"

    # 5. Pending Recon
    assert classify_rfd("Client claims already settled long ago, pending recon with bank") == "Pending Recon"

    # 6. Informant speaker guardrail (e.g. caretaker Bong Bong, store staff, HOA guard)
    c1 = classify_contact("", "", "Positive address, unit negative; per caretaker Bong Bong, property confirmed owned by client.")
    assert c1.category_label == "Informant"

    c2 = classify_contact("", "", "Address verified positive; only store staff present. Per staff, client out of town.")
    assert c2.category_label == "Informant"

    c3 = classify_contact("", "", "House closed; per HOA guard, client still resides in subdivision.")
    assert c3.category_label == "Informant"


def test_remarks_lab_test_mode_vs_real_mode():
    """Verify that test mode grabs from MESSAGE while real mode grabs from FINAL REMARKS."""
    df = pd.DataFrame({
        "Account Number": ["10101"],
        "CH Code": ["CH101"],
        "Contact Person": ["Juana Cruz"],
        "Contact Relation": ["Asawa"],
        "MESSAGE": ["Spoke to asawa Juana, unit in garage PTP tomorrow"],
        "FINAL REMARKS": ["ECA AUTO_S.P. MADRID_HOME_09/18/2026 PTP on Monday confirmed"],
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="FIELD RSULT", index=False)
    data = buf.getvalue()

    pipeline = RemarksLabPipeline()

    # 1. Test Mode = True -> raw remark must come from MESSAGE
    report_test = pipeline.process_file(io.BytesIO(data), test_mode=True)
    assert report_test.test_mode is True
    assert "Spoke to asawa Juana" in report_test.rows[0].raw_remarks

    # 2. Test Mode = False (Real Mode) -> raw remark must come from FINAL REMARKS
    report_real = pipeline.process_file(io.BytesIO(data), test_mode=False)
    assert report_real.test_mode is False
    assert "PTP on Monday confirmed" in report_real.rows[0].raw_remarks


def test_remarks_lab_no_user_name_in_summarized_remark():
    """Verify that contact person / user's name is NOT included in trimmed_statement."""
    df = pd.DataFrame({
        "Account Number": ["20202"],
        "CH Code": ["CH202"],
        "Contact Person": ["JUAN DELA CRUZ JR."],
        "Contact Relation": ["Self"],
        "FINAL REMARKS": [
            "Borrower confirms unit is at registered address, promising settlement next week."
        ],
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="FIELD RSULT", index=False)
    data = buf.getvalue()

    pipeline = RemarksLabPipeline()
    report = pipeline.process_file(io.BytesIO(data), test_mode=False)

    trimmed = report.rows[0].trimmed_statement
    # Contact person name must NOT be prefixed or included in the final trimmed statement
    assert "JUAN DELA CRUZ JR." not in trimmed
    assert "JUAN" not in trimmed


def test_remarks_lab_web_test_mode_toggle_handling():
    """Verify web endpoint handles test_mode toggle form submission."""
    app = create_demo_app()
    client = TestClient(app)

    df = pd.DataFrame({
        "Account Number": ["30303"],
        "Contact Person": ["Pedro"],
        "MESSAGE": ["Raw collector message text"],
        "FINAL REMARKS": ["Production final remarks text"],
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="FIELD RSULT", index=False)
    data = buf.getvalue()

    excel_mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    # POST with test_mode=true
    resp_test = client.post(
        "/remarks-lab/analyze",
        data={"test_mode": "true"},
        files={"workbook": ("test.xlsx", data, excel_mime)},
    )
    assert resp_test.status_code == 200
    assert "TEST (Source: MESSAGE)" in resp_test.text
    assert "Raw collector message text" in resp_test.text

    # POST without test_mode (Real Mode)
    resp_real = client.post(
        "/remarks-lab/analyze",
        data={},
        files={"workbook": ("test.xlsx", data, excel_mime)},
    )
    assert resp_real.status_code == 200
    assert "REAL (Source: FINAL REMARKS)" in resp_real.text
    assert "Production final remarks text" in resp_real.text


def test_strip_user_name_from_raw_remark_with_eca_header():
    """Verify that borrower/contact name is stripped from raw remark containing ECA header."""
    raw_text = (
        "ECA AUTO_S.P. MADRID_HOME_09/02/2026 LEYSON, JEANIE BARRACA - "
        "neg, according to the agent the account has been resolved"
    )

    # 1. Direct trimmer test with contact person provided
    res1 = trim_and_format(raw_text, contact_person="LEYSON, JEANIE BARRACA", strip_name=True)
    assert "LEYSON, JEANIE BARRACA" not in res1.trimmed_text
    assert res1.trimmed_text == (
        "ECA AUTO_S.P. MADRID_HOME_09/02/2026 neg, according to the agent the account has been resolved"
    )

    # 2. Direct trimmer test with empty contact person (regex fallback)
    res2 = trim_and_format(raw_text, contact_person="", strip_name=True)
    assert "LEYSON, JEANIE BARRACA" not in res2.trimmed_text
    assert res2.trimmed_text == (
        "ECA AUTO_S.P. MADRID_HOME_09/02/2026 neg, according to the agent the account has been resolved"
    )

    # 3. Full pipeline test
    df = pd.DataFrame({
        "Account Number": ["40404"],
        "CH Code": ["CH404"],
        "Contact Person": ["LEYSON, JEANIE BARRACA"],
        "Contact Relation": ["Self"],
        "FINAL REMARKS": [raw_text],
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="FIELD RSULT", index=False)

    pipeline = RemarksLabPipeline()
    report = pipeline.process_file(io.BytesIO(buf.getvalue()), test_mode=False)

    statement = report.rows[0].trimmed_statement
    assert "LEYSON, JEANIE BARRACA" not in statement
    assert "LEYSON" not in statement
    assert statement == (
        "ECA AUTO_S.P. MADRID_HOME_09/02/2026 neg, according to the agent the account has been resolved"
    )


def test_cut_eca_header_when_sending_to_llm_and_remember_per_remark():
    """Verify that ECA header is cut from prompt to save tokens and remembered dynamically per remark."""
    from mc03.services.remarks_lab.groq_summarizer import GroqRemarksSummarizer

    summarizer = GroqRemarksSummarizer(api_key="gsk_test_mock")

    # 1. Verify _build_prompt cuts out the ECA header prefix
    raw1 = "ECA AUTO_S.P. MADRID_HOME_09/02/2026 neg, according to the agent the account has been resolved"
    messages = summarizer._build_prompt(raw1, max_chars=140, is_retry=False)
    user_content = messages[1]["content"]

    # Must NOT contain ECA header in prompt sent to LLM
    assert "ECA AUTO_S.P. MADRID_HOME_09/02/2026" not in user_content
    assert "neg, according to the agent the account has been resolved" in user_content

    # 2. Test with a different remark and different ECA header format (date with hyphen, office location)
    raw2 = "ECA S.P. MADRID_OFFICE_10/05/2026: borrower visited office and promised full payment tomorrow"
    messages2 = summarizer._build_prompt(raw2, max_chars=140, is_retry=False)
    user_content2 = messages2[1]["content"]

    assert "ECA S.P. MADRID_OFFICE_10/05/2026" not in user_content2
    assert "borrower visited office and promised full payment tomorrow" in user_content2

    # 3. Test that mock LLM provider returns summary with the remembered dynamic header re-attached
    long_raw = (
        "ECA AUTO_S.P. MADRID_HOME_09/02/2026 neg, according to the field collection agent the account has been resolved "
        "and client has paid in full at the branch according to the receipt shown to the field officer during the second visit."
    )
    def mock_llm_provider(prompt_text: str, max_chars: int = 140, is_retry_shorter: bool = False) -> str:
        # Verify prompt received by mock LLM has NO ECA header
        assert "ECA" not in prompt_text
        return "neg per agent, account resolved in full at branch"

    res = trim_and_format(long_raw, llm_provider=mock_llm_provider, strip_name=True)
    assert res.trimmed_text.startswith("ECA AUTO_S.P. MADRID_HOME_09/02/2026")
    assert "neg per agent, account resolved in full at branch" in res.trimmed_text


def test_no_okay_na_to_in_ui_or_export():
    """Verify that 'OKAY NA TO' is completely removed from UI template and exports."""
    app = create_demo_app()
    client = TestClient(app)

    df = pd.DataFrame({
        "Account Number": ["50505"],
        "CH Code": ["CH505"],
        "Contact Person": ["Ana Ramos"],
        "Contact Relation": ["Self"],
        "FINAL REMARKS": ["Confirmed payment made"],
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="FIELD RSULT", index=False)
    data = buf.getvalue()

    # 1. UI response
    resp = client.post(
        "/remarks-lab/analyze",
        data={},
        files={"workbook": ("test.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    assert "OKAY NA TO" not in resp.text

    # 2. Excel export
    resp_exp = client.post(
        "/remarks-lab/export",
        files={"workbook": ("test.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp_exp.status_code == 200
    df_exp = pd.read_excel(io.BytesIO(resp_exp.content), sheet_name="FIELD RESULT (EVALUATED)")
    assert "char checker" in df_exp.columns
    # Check that OKAY NA TO is not in char checker
    assert "OKAY NA TO" not in list(df_exp["char checker"].fillna(""))


def test_relation_classifier_remark_overrides_metadata_conflict():
    """Verify remark text wins when it conflicts with metadata, and label indicates both checked."""
    # Metadata says 'Mother', but remark clearly states collector spoke to 'neighbor'
    res = classify_contact(
        contact_person="Maria Santos",
        contact_relation="Mother",
        statement="ECA AUTO_S.P. MADRID_HOME_09/02/2026 Spoke to neighbor ms delia, client moved to Laguna"
    )
    assert res.category_label == "Informant"
    assert res.normalized_role == "Neighbor"
    assert "remark:neighbor" in res.matched_keyword
    assert "contact field listed (mother)" in res.matched_keyword


def test_relation_classifier_remark_no_contact_overrides_metadata():
    """Verify remark indicating house closed / nobody around overrides metadata contact listing."""
    res = classify_contact(
        contact_person="Juan Dela Cruz",
        contact_relation="Self",
        statement="House closed gate locked, nobody around to answer"
    )
    assert res.category_label == "none"
    assert res.normalized_role == "none"


def test_relation_classifier_agreement_both_sources():
    """Verify agreement label when both remark text and metadata point to the same role."""
    res = classify_contact(
        contact_person="Corazon Santos",
        contact_relation="Mother",
        statement="Spoke with mother, client promises settlement on Friday"
    )
    assert res.category_label == "Representative"
    assert res.normalized_role == "Mother"
    assert "remark:mother" in res.matched_keyword
    assert "also listed in contact field" in res.matched_keyword


def test_trim_method_tracking_and_counts():
    """Verify trim_method ('none' | 'llm' | 'fallback') is tracked explicitly and counts match."""
    from mc03.services.remarks_lab.groq_summarizer import GroqRemarksSummarizer

    # 1. Row that fits within 200 chars as-is
    res_none = trim_and_format("Short remark fits easily", strip_name=True)
    assert res_none.trim_method == "none"

    # 2. Row trimmed by LLM
    long_remark = (
        "ECA AUTO_S.P. MADRID_HOME_09/02/2026 Upon visit to the given residential address, the house was open "
        "and collector was able to interview the client who stated that due to medical emergency and delayed salary "
        "from company they were unable to pay on time, promising to settle full account next Friday at branch office."
    )
    def mock_llm(text: str, max_chars: int = 140, is_retry_shorter: bool = False):
        return "Spoke with client, promised settlement on Monday"

    res_llm = trim_and_format(long_remark, llm_provider=mock_llm, strip_name=True)
    assert res_llm.trim_method == "llm"

    # 3. Row trimmed by fallback (no LLM provider)
    res_fb = trim_and_format(long_remark, llm_provider=None, strip_name=True)
    assert res_fb.trim_method == "fallback"

    # 4. Pipeline execution with offline summarizer: 1 untrimmed, 1 fallback trimmed
    df = pd.DataFrame({
        "Account Number": ["60601", "60602"],
        "CH Code": ["CH01", "CH02"],
        "Contact Person": ["A", "B"],
        "Contact Relation": ["Self", "Self"],
        "FINAL REMARKS": [
            "Short remark",
            long_remark
        ],
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="FIELD RSULT", index=False)

    offline_summarizer = GroqRemarksSummarizer(api_key="")
    pipeline = RemarksLabPipeline(summarizer=offline_summarizer)  # Offline -> fallback used for long remark
    report = pipeline.process_file(io.BytesIO(buf.getvalue()))

    assert report.total_rows == 2
    assert report.char_overflow_prevented_count == 1
    assert report.ai_summarized_count == 0
    assert report.rule_based_fallback_count == 1
    assert report.untrimmed_count == 1
    # Trimmed rows must split cleanly into AI vs Rule Fallback
    assert report.ai_summarized_count + report.rule_based_fallback_count == report.char_overflow_prevented_count
    assert report.rows[0].trim_method == "none"
    assert report.rows[1].trim_method == "fallback"


def test_filter_ai_summarized_vs_rule_based_ui():
    """Verify Remarks Lab UI renders AI-Summarized and Rule-Based Fallback filter pills and row data-trim-method."""
    app = create_demo_app()
    client = TestClient(app)

    long_remark = (
        "ECA AUTO_S.P. MADRID_HOME_09/02/2026 Upon visit to the given residential address, the house was open "
        "and collector was able to interview the client who stated that due to medical emergency and delayed salary "
        "from company they were unable to pay on time, promising to settle full account next Friday at branch office."
    )
    df = pd.DataFrame({
        "Account Number": ["70701", "70702"],
        "Contact Person": ["Juana", "Maria"],
        "FINAL REMARKS": [
            "Short remark",
            long_remark,
        ],
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="FIELD RSULT", index=False)

    resp = client.post(
        "/remarks-lab/analyze",
        data={},
        files={"workbook": ("test.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    assert "pill-ai-summarized" in resp.text
    assert "pill-rule-based-fallback" in resp.text
    assert 'data-trim-method="none"' in resp.text
    assert ('data-trim-method="fallback"' in resp.text or 'data-trim-method="llm"' in resp.text)
    assert ("RULE FALLBACK" in resp.text or "AI SUMMARIZED" in resp.text)


def test_csu_column_mapping_column_j_not_status_column_o():
    """Verify that 'COLLECTION STATUS UPDATE\\nAGENT/ADMIN' (Column J) maps to manual_csu, NOT 'status' (Column O)."""
    from mc03.services.remarks_lab.parser import parse_field_result_sheet

    df = pd.DataFrame({
        "ACCT NUM": ["10001"],
        "FINAL REMARKS": ["Client promised payment on Monday"],
        "COLLECTION STATUS UPDATE\nAGENT/ADMIN": ["CLIENT POSITIVE/UNIT POSITIVE (WITH Actual Contact - Both)"],
        "status": ["REPO"],  # Column O in RCBC sheets
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="FIELD RSULT", index=False)

    parsed = parse_field_result_sheet(io.BytesIO(buf.getvalue()))
    assert len(parsed) == 1
    # manual_csu MUST be from Column J, NOT Column O ('REPO')
    assert parsed[0].manual_csu == "CLIENT POSITIVE/UNIT POSITIVE (WITH Actual Contact - Both)"
    assert parsed[0].manual_csu != "REPO"


def test_remarks_lab_prompt_finetuning_endpoint():
    """Test POST /remarks-lab/finetune-prompt with empty and valid feedback."""
    app = create_demo_app()
    client = TestClient(app)

    # 1. Reject empty items
    res_empty = client.post("/remarks-lab/finetune-prompt", json={"items": []})
    assert res_empty.status_code == 400
    assert "No feedback provided" in res_empty.json()["error"]

    # 2. Reject items with blank feedback strings
    res_blank = client.post("/remarks-lab/finetune-prompt", json={"items": [{"feedback": "   "}]})
    assert res_blank.status_code == 400

    # 3. Accept valid feedback and return diagnosis + prompt
    payload = {
        "items": [
            {
                "row_index": 12,
                "account_number": "ACC-101",
                "raw_remarks": "Met borrower wife Mrs. Santos at house. Unit is in garage.",
                "cleaned_remarks": "Met borrower wife Mrs. Santos, unit in garage.",
                "feedback": "Do not include the person's last name Santos. Keep it strictly anonymous."
            },
            {
                "row_index": 15,
                "account_number": "ACC-102",
                "raw_remarks": "Borrower promised to pay 5,000 on September 25 via GCash.",
                "cleaned_remarks": "Borrower will settle amount soon.",
                "feedback": "Preserve the exact amount 5000 and date Sept 25."
            }
        ]
    }
    res_valid = client.post("/remarks-lab/finetune-prompt", json=payload)
    assert res_valid.status_code == 200
    data = res_valid.json()
    assert data["total_feedback_items"] == 2
    assert len(data["items_processed"]) == 2
    assert "ai_analysis" in data
    assert len(data["ai_analysis"]) > 50


def test_heuristic_finetune_row_and_account_formatting():
    """Verify that heuristic fallback clearly formats actual Row # and Account Number separately."""
    from mc03.services.web import _generate_heuristic_finetune

    items = [
        {
            "row_index": 20133,
            "account_number": "80502100885",
            "raw_remarks": "Met ch brother.",
            "cleaned_remarks": "Met ch brother.",
            "feedback": "use REPRESENTATIVE REFUSED TO DISCLOSE RFD if theres no interaction"
        },
        {
            "account_number": "999999",
            "feedback": "Shorten to 100 chars."
        }
    ]
    analysis = _generate_heuristic_finetune(items, "current prompt", api_error="No module named 'groq'")
    assert "Row #20133 (Acct: 80502100885):" in analysis
    assert "Row #2 (Acct: 999999):" in analysis
    assert "Groq API error: No module named 'groq'" in analysis


def test_official_csu_and_rfd_lists_in_summarizer():
    """Verify that TwoTierRemarksSummarizer uses the official 42 CSUs and 27 RFDs and strictly excludes 'Unit Impounded'."""
    from engine.summarizer import (
        TwoTierRemarksSummarizer,
        OFFICIAL_RCBC_CSU_OPTIONS,
        OFFICIAL_RCBC_RFD_OPTIONS,
    )
    from mc03.services.remarks_lab.pipeline import _is_rfd_match

    # Verify counts
    assert len(OFFICIAL_RCBC_CSU_OPTIONS) == 42
    assert len(OFFICIAL_RCBC_RFD_OPTIONS) == 27

    # Verify 'Unit Impounded' is NOT an RFD, but 'LTO APPREHENSION/NO ORCR/HPG' is
    assert "Unit Impounded" not in OFFICIAL_RCBC_RFD_OPTIONS
    assert "UNIT IMPOUNDED" not in OFFICIAL_RCBC_RFD_OPTIONS
    assert "LTO APPREHENSION/NO ORCR/HPG" in OFFICIAL_RCBC_RFD_OPTIONS

    # Verify key new categories are present
    assert "BANK ACCOUNT ON-HOLD/UNDER GARNISHMENT" in OFFICIAL_RCBC_RFD_OPTIONS
    assert "NO CLIENT/ REPRESENTATIVE" in OFFICIAL_RCBC_RFD_OPTIONS
    assert "NEGATIVE CLIENT / NEGATIVE UNIT (REAL / HARD SKIPS)" in OFFICIAL_RCBC_CSU_OPTIONS

    # Verify prompt builder contains the official lists and guidelines
    summarizer = TwoTierRemarksSummarizer(api_key="sk-test")
    messages = summarizer._build_prompt("Vehicle was impounded by LTO; coordinate with lawyer", contact_person="Maria")
    prompt_content = messages[0]["content"]

    assert "ALLOWED_CSU (Select EXACTLY one verbatim from this official RCBC list):" in prompt_content
    assert "ALLOWED_RFD (Select EXACTLY one verbatim from this official RCBC list):" in prompt_content
    assert "LTO APPREHENSION/NO ORCR/HPG" in prompt_content
    assert "NEVER output 'Unit Impounded'" in prompt_content

    # Verify matching logic matches impounded to LTO
    assert _is_rfd_match("LTO APPREHENSION/NO ORCR/HPG", "Unit impounded") is True


def test_remarks_lab_ui_bank_upload_in_col4_and_feedback_in_col6():
    """Verify bank upload content is moved to Col 4 with length, and Col 6 has feedback textarea."""
    app = create_demo_app()
    client = TestClient(app)

    df = pd.DataFrame({
        "Account Number": ["ACC-888"],
        "Contact Person": ["Maria"],
        "Contact Relation": ["Ate"],
        "Remarks": ["PTP on Friday next week for settlement"]
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="FIELD RSULT", index=False)

    resp = client.post(
        "/remarks-lab/analyze",
        files={"workbook": ("test.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200
    html = resp.text

    # Header check
    assert "Raw &amp; Cleaned Remarks (Bank Upload)" in html or "Raw & Cleaned Remarks (Bank Upload)" in html
    assert "Feedback for Fine-Tuning LLM Prompt" in html

    # Col 4 check
    assert "Bank Upload Cleaned:" in html
    assert "cleaned-remark-text" in html
    assert "Len: " in html
    assert "/200" in html

    # Col 6 check
    assert "prompt-feedback-input" in html
    assert "Leave blank (default) or type feedback for AI prompt tuning..." in html
    assert "Default: blank" in html

    # Modal and button check
    assert "finetuneModal" in html
    assert "sendFeedbackBtn" in html
    assert "Send Feedback to AI" in html
    assert "applyRulesBtn" in html
    assert "Apply to Active Models" in html


def test_remarks_lab_csv_import_and_test_mode_agreement():
    """Verify exported Test Mode CSV can be re-imported, correctly mapping Original CSU/RFD with 0 LLM usage."""
    # Create CSV mimicking exported test format
    csv_text = (
        '"Row Index","Account Number","CH Code","Contact Person","Contact Relation","Category Label","Role","Raw Remark","Original CSU","Original RFD","COLLECTION STATUS UPDATE","RFD","DETAILED RFD","FINAL REMARKS","Char Count","char checker"\r\n'
        '"20133","ACC-001","CH1","Jane Doe","Wife","Representative","Wife","as per ch wife unit not seen ch not around ptp on Friday","CLIENT POSITIVE/UNIT NEGATIVE (WITHOUT Actual Contact)","REPRESENTATIVE REFUSED TO DISCLOSE RFD","OLD CSU","OLD RFD","OLD DETAILED","ECA AUTO_TEST_01/01/2026 as per ch wife unit not seen ch not around ptp on Friday","85","OK"\r\n'
        '"20134","ACC-002","CH2","Self","Borrower","Borrower","Self","ch promised to settle full balance on Monday","CLIENT POSITIVE/UNIT POSITIVE (WITH Actual Contact)","Financial Difficulty","OLD CSU","OLD RFD","OLD DETAILED","ECA AUTO_TEST_01/01/2026 ch promised to settle full balance on Monday","71","OK"\r\n'
    )
    csv_bytes = csv_text.encode("utf-8")

    # 1. Verify parser mapping
    raw_rows = parse_field_result_sheet(csv_bytes, date_from=None, date_to=None)
    assert len(raw_rows) == 2
    assert raw_rows[0].row_index == 20133
    assert raw_rows[0].account_number == "ACC-001"
    assert raw_rows[0].contact_relation == "Wife"
    assert raw_rows[0].manual_csu == "CLIENT POSITIVE/UNIT NEGATIVE (WITHOUT Actual Contact)"
    assert raw_rows[0].manual_rfd == "REPRESENTATIVE REFUSED TO DISCLOSE RFD"
    assert "as per ch wife" in raw_rows[0].raw_remarks

    # 2. Verify pipeline execution in test mode (bypass LLM)
    mock_llm_called = False
    class DummySummarizer:
        is_available = True
        model = "dummy"
        batch_processing = False
        def summarize(self, text, max_chars=200):
            nonlocal mock_llm_called
            mock_llm_called = True
            return text[:max_chars]

    pipeline = RemarksLabPipeline(summarizer=DummySummarizer())
    report = pipeline.process_file(csv_bytes, test_mode=True)
    assert not mock_llm_called, "LLM summarizer must not be called during test mode evaluation"
    assert report.total_rows == 2
    assert report.test_mode is True
    # Row 20133 should be an exact match
    row1 = report.rows[0]
    assert row1.row_index == 20133
    assert row1.manual_csu == "CLIENT POSITIVE/UNIT NEGATIVE (WITHOUT Actual Contact)"
    assert row1.predicted_csu == "CLIENT POSITIVE/UNIT NEGATIVE (WITHOUT Actual Contact)"
    assert row1.csu_match is True

    # 3. Verify web endpoint accepts CSV upload
    app = create_demo_app()
    client = TestClient(app)
    resp = client.post(
        "/remarks-lab/analyze",
        files={"workbook": ("RCBC_FIELD_RESULT_TEST_EVALUATED.csv", csv_bytes, "text/csv")},
        data={"test_mode": "true"},
    )
    assert resp.status_code == 200
    html = resp.text
    assert "20133" in html
    assert "ACC-001" in html
    assert "Representative" in html


def test_csv_with_collection_status_update_and_date_filter():
    """Verify that 'COLLECTION STATUS UPDATE' column is not mistaken for a date column,
    and date filtering works properly via remarks embedded date or falls back gracefully."""
    from datetime import date
    csv_content = (
        '"Row Index","Account Number","CH Code","Contact Person","Contact Relation","Category Label","Role","Raw Remark","Original CSU","Original RFD","COLLECTION STATUS UPDATE","RFD","DETAILED RFD","FINAL REMARKS","Char Count","char checker"\n'
        '"20129","80502100885","BKAL-1","LEYSON, JEANIE","none stated","none","","neg, resolved","CLIENT POSITIVE/UNIT POSITIVE (WITHOUT Actual Contact - Client)","REPRESENTATIVE REFUSED TO DISCLOSE RFD","CLIENT POSITIVE/UNIT NEGATIVE (WITHOUT Actual Contact)","NO CLIENT / REPRESENTATIVE REACHED","Details","ECA AUTO_S.P. MADRID_HOME_09/02/2026 neg, resolved","94",""\n'
        '"20130","1000069516","BKAL-2","NUGUID, DORINA","none stated","none","","Done repo unit","CLIENT POSITIVE/UNIT POSITIVE (WITH Actual Contact - Both)","BORROWER REFUSED TO DISCLOSE RFD","CLIENT NEGATIVE/UNIT NEGATIVE (FOR FURTHER VISIT/PROBING)","NO CLIENT / REPRESENTATIVE REACHED","Details","ECA AUTO_S.P. MADRID_HOME_09/02/2026 Done repo unit","63",""\n'
    )
    csv_bytes = csv_content.encode("utf-8")

    pipeline = RemarksLabPipeline()
    # 1. Date window matching 2026-09-02 should return both rows
    report = pipeline.process_file(
        io.BytesIO(csv_bytes),
        date_from=date(2026, 9, 2),
        date_to=date(2026, 9, 2),
        test_mode=True,
    )
    assert report.total_rows == 2
    assert report.rows[0].row_date == "2026-09-02"
    assert report.rows[0].account_number == "80502100885"

    # 2. Date window not matching (e.g. 2026-09-05) should filter out the rows
    report_empty = pipeline.process_file(
        io.BytesIO(csv_bytes),
        date_from=date(2026, 9, 5),
        date_to=date(2026, 9, 5),
        test_mode=True,
    )
    assert report_empty.total_rows == 0

    # 3. HTTP POST to /remarks-lab/analyze with date filter
    app = create_demo_app()
    client = TestClient(app)
    resp = client.post(
        "/remarks-lab/analyze",
        files={"workbook": ("test_export.csv", csv_bytes, "text/csv")},
        data={"date_from": "2026-09-02", "date_to": "2026-09-02", "test_mode": "true"},
    )
    assert resp.status_code == 200
    assert "80502100885" in resp.text
    assert "1000069516" in resp.text
    assert "Active Records: <strong class=\"font-mono text-indigo-800\">2</strong> rows" in resp.text


def test_extract_structured_rules_from_feedback():
    """Verify _extract_structured_rules derives actionable CSU/RFD and summary rules."""
    from mc03.services.web import _extract_structured_rules

    items = [
        {"feedback": "use REPRESENTATIVE REFUSED TO DISCLOSE RFD if theres no interaction"},
        {"feedback": "vehicle was impounded by LTO and should not show unit impounded"},
        {"feedback": "client moved out to another province according to informant"},
        {"feedback": "leave rfd empty since client is unknown in area"},
        {"feedback": "keep summary note strictly short under 120 chars"},
        {"feedback": "never include borrower names in output"},
    ]

    rules = _extract_structured_rules(items)
    csu_rfd = rules["csu_rfd_rules"]
    summary = rules["summary_rules"]

    # Verify CSU/RFD extraction
    assert any("REPRESENTATIVE REFUSED TO DISCLOSE RFD" in r for r in csu_rfd)
    assert any("LTO APPREHENSION/NO ORCR/HPG" in r or "impounded" in r.lower() for r in csu_rfd)
    assert any("MOVED OUT" in r for r in csu_rfd)
    assert any("empty" in r.lower() for r in csu_rfd)

    # Verify summary rules extraction
    assert any("100-120" in r or "short" in r.lower() for r in summary)
    assert any("borrower" in r.lower() or "names" in r.lower() for r in summary)


def test_remarks_lab_apply_rules_endpoint_and_summarizer_injection():
    """Verify POST /remarks-lab/apply-rules persists rules and TwoTierRemarksSummarizer immediately injects them into prompt."""
    import json
    from engine.summarizer import TwoTierRemarksSummarizer, CUSTOM_RULES_FILE

    app = create_demo_app()
    client = TestClient(app)

    test_rules = {
        "csu_rfd_rules": [
            "Use REPRESENTATIVE REFUSED TO DISCLOSE RFD if there was no direct borrower contact",
            "Never classify as Unit Impounded",
        ],
        "summary_rules": [
            "Cap summary strictly at 110 characters",
        ],
    }

    resp = client.post("/remarks-lab/apply-rules", json=test_rules)
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "success"
    assert "applied_rules" in data

    # Verify file content
    assert CUSTOM_RULES_FILE.exists()
    with open(CUSTOM_RULES_FILE, "r", encoding="utf-8") as f:
        stored = json.load(f)
    assert stored["csu_rfd_rules"] == test_rules["csu_rfd_rules"]
    assert stored["summary_rules"] == test_rules["summary_rules"]

    # Verify prompt injection in TwoTierRemarksSummarizer
    summarizer = TwoTierRemarksSummarizer(api_key="sk-test")
    messages = summarizer._build_prompt("Vehicle seen parked in garage, spoke with wife")
    prompt = messages[0]["content"]

    assert "### ACTIVE HUMAN REVIEWER TUNED DIRECTIVES (HIGHEST PRIORITY):" in prompt
    assert "Use REPRESENTATIVE REFUSED TO DISCLOSE RFD if there was no direct borrower contact" in prompt
    assert "Never classify as Unit Impounded" in prompt
    assert "Cap summary strictly at 110 characters" in prompt

    # Clean up by posting empty rules
    client.post("/remarks-lab/apply-rules", json={"csu_rfd_rules": [], "summary_rules": []})
    messages_clean = summarizer._build_prompt("Vehicle seen parked in garage")
    assert "ACTIVE HUMAN REVIEWER TUNED DIRECTIVES" not in messages_clean[0]["content"]


def test_empty_rfd_classification_and_pipeline_preservation():
    """Verify empty RFD classification for unknown clients and preservation in pipeline."""
    from mc03.services.remarks_lab.ranker import classify_rfd, RFD_NO_REACH, RFD_MOVED_OUT, RFD_BORROWER_REFUSED
    from mc03.services.remarks_lab.pipeline import RemarksLabPipeline

    # 1. Informant unknown / zero info gathered -> ""
    assert classify_rfd("Informant confirmed client is unknown in the area; zero info gathered", "") == ""
    assert classify_rfd("Informant stated they don't know the client at all", "") == ""

    # 2. Gate closed / padlocked / house closed (still resides) -> NO CLIENT / REPRESENTATIVE REACHED
    assert classify_rfd("Gate closed, padlocked, client not around during visit", "") == RFD_NO_REACH
    assert classify_rfd("House closed, client is out for work", "") == RFD_NO_REACH

    # 3. Explicit moved out -> Moved Out
    assert classify_rfd("Informant confirmed borrower moved out to Cavite", "") == RFD_MOVED_OUT

    # 4. Borrower contact -> BORROWER REFUSED TO DISCLOSE RFD
    assert classify_rfd("Spoke with borrower directly, promise to settle next week", "Borrower") == RFD_BORROWER_REFUSED

    # 5. Verify pipeline preserves explicit empty string from rule and LLM
    pipeline = RemarksLabPipeline(summarizer=None)
    csv_text = (
        "Row Index,Account Number,Remarks\n"
        "1,ACC-001,Informant confirmed client is unknown in area\n"
    )
    report = pipeline.process_file(io.BytesIO(csv_text.encode("utf-8")), test_mode=True)
    assert report.total_rows == 1
    assert report.rows[0].predicted_rfd == ""


def test_csu_pending_recon_does_not_mismatch_positive_unit():
    """Verify that PENDING RECON does not falsely match CLIENT POSITIVE/UNIT POSITIVE entries."""
    from mc03.services.remarks_lab.pipeline import _is_csu_match

    # Must NOT match when one is Pending Recon and the other is a contact/unit matrix status
    assert not _is_csu_match(
        "PENDING RECON",
        "CLIENT POSITIVE/UNIT POSITIVE (WITHOUT Actual Contact - Client)"
    )
    assert not _is_csu_match(
        "CLIENT POSITIVE/UNIT POSITIVE (WITHOUT Actual Contact - Client)",
        "PENDING RECON"
    )

    # Must match when both are recon
    assert _is_csu_match("PENDING RECON", "Pending Recon")
    assert _is_csu_match("Pending Recon", "Recon")


def test_csu_and_rfd_decision_reasoning_pipeline_and_ui():
    """Verify CSU and RFD decision reasoning are populated and rendered in Column 5."""
    from mc03.services.remarks_lab.pipeline import RemarksLabPipeline
    from mc03.services.remarks_lab.ranker import rank_rfd_and_csu

    # 1. Test deterministic reasoning generation
    res = rank_rfd_and_csu(
        "Spoke with sister, borrower at work, unit not seen",
        "Spoke with sister, borrower at work, unit not seen",
        "Representative",
        "Sister",
    )
    assert res.csu_reasoning != ""
    assert res.rfd_reasoning != ""
    assert "representative" in res.csu_reasoning.lower() or "not borrower" in res.csu_reasoning.lower()

    # 2. Test pipeline population
    pipeline = RemarksLabPipeline(summarizer=None)
    csv_text = (
        "Row Index,Account Number,Contact Relation,Remarks\n"
        "1,ACC-101,Sister,Spoke with sister borrower at work unit not seen\n"
    )
    report = pipeline.process_file(io.BytesIO(csv_text.encode("utf-8")), test_mode=True)
    assert len(report.rows) == 1
    row = report.rows[0]
    assert row.csu_reasoning != ""
    assert row.rfd_reasoning != ""

    # 3. Test UI rendering via /remarks-lab/analyze
    app = create_demo_app()
    client = TestClient(app)
    resp = client.post(
        "/remarks-lab/analyze",
        files={"workbook": ("test.csv", csv_text.encode("utf-8"), "text/csv")},
        data={"test_mode": "true"},
    )
    assert resp.status_code == 200
    assert "Why" in resp.text
    assert row.csu_reasoning in resp.text
    assert row.rfd_reasoning in resp.text


def test_filter_out_blank_or_na_account_numbers():
    """Verify that rows with blank or N/A account numbers are not loaded into Remarks Lab."""
    from mc03.services.remarks_lab.parser import parse_field_result_sheet

    csv_text = (
        "Row Index,Account Number,Remarks\n"
        "1,1000101,Valid account row 1\n"
        "2,N/A,Unit repo scan without account\n"
        "3,,Blank account row\n"
        "4,na,Lowercase na account\n"
        "5,None,None account\n"
        "6,-,Dash account\n"
        "7,1000102,Valid account row 2\n"
        "8,  ,Whitespace account\n"
    )
    parsed = parse_field_result_sheet(csv_text.encode("utf-8"), test_mode=True)
    assert len(parsed) == 2
    assert [r.account_number for r in parsed] == ["1000101", "1000102"]





